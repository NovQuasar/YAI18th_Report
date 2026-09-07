"""
YAI 보고서 제출 현황 → 엑셀 변환 스크립트
------------------------------------------
results/ 폴더의 평가 결과(JSON)를 모아 need.md 형식의 엑셀 파일로 만든다.
"개인 보고서"는 학회원 개인별, "팀 Discussion"은 팀별로 집계한다.

사용법:
  conda activate yai
  python export_excel.py                              # 개인 보고서 (기본값)
  python export_excel.py --type "팀 Discussion"        # 팀 Discussion
  python export_excel.py --roster members.txt          # 미제출자까지 표시하려면 명단 파일 지정
  python export_excel.py --output 어딘가/파일명.xlsx     # 출력 경로 지정

전제:
  - app.py에서 평가할 때 "학회원 이름"과 "주차"를 입력해야 해당 기록이 집계된다.
    (두 값이 비어 있는 과거 기록은 건너뛰고 경고만 출력한다.) "팀 Discussion"은 학회원 이름 대신
    "팀 이름"과 주차가 채워져 있어야 한다.
  - --roster로 전체 명단 파일(한 줄에 하나, '#'로 시작하면 주석)을 넘기면 "❌ 미제출" 행도 채워진다.
    "개인 보고서"는 학회원 이름 명단을, "팀 Discussion"은 팀 이름 명단을 넘겨야 한다. 넘기지 않으면
    그 행은 비워둔다.
"""

import argparse
import json
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).parent
RESULTS_DIR = BASE_DIR / "results"

# 보고서 종류별 집계 설정.
# - entity: "member"면 학회원(member_name) 기준, "team"이면 팀(team_name) 기준으로 묶는다.
# - criteria: app.py REPORT_TYPES[report_type]["criteria"]와 동일한 (JSON 키, 엑셀 열 라벨) 목록.
# - has_team_type: eval_criteria_personal.py만 "team_type"(강의팀/논문팀)을 채점 결과에 포함한다.
REPORT_CONFIGS = {
    "개인 보고서": {
        "entity": "member",
        "criteria": [
            ("이해도", "이해도(5)"),
            ("가독성", "가독성(5)"),
            ("시각자료", "시각자료(3)"),
            ("토론", "토론(3)"),
        ],
        "has_team_type": True,
    },
    "팀 Discussion": {
        "entity": "team",
        "criteria": [
            ("팀세션_진행내용", "팀세션_진행내용(5)"),
            ("Discussion_과정", "Discussion_과정(5)"),
            ("Discussion_확장성", "Discussion_확장성(2)"),
        ],
        "has_team_type": False,
    },
}

HEADER_FILL = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
HEADER_FONT = Font(bold=True)


def _team_name_of(record):
    return (record.get("team_name") or record["result"].get("team_name") or "").strip()


def load_records(report_type):
    """results/*.json 중 지정한 report_type + (entity에 맞는 이름)/주차가 채워진 기록만 로드."""
    config = REPORT_CONFIGS[report_type]
    entity = config["entity"]
    records = []
    skipped = 0
    for path in sorted(RESULTS_DIR.glob("*.json")):
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except (json.JSONDecodeError, OSError):
            continue
        if data.get("report_type") != report_type:
            continue
        week = data.get("week")
        key = (data.get("member_name") or "").strip() if entity == "member" else _team_name_of(data)
        if not key or not week:
            skipped += 1
            continue
        data["_entity_key"] = key
        records.append(data)
    if skipped:
        entity_label = "학회원 이름" if entity == "member" else "팀 이름"
        print(f"[안내] {entity_label}/주차가 비어 있어 건너뛴 '{report_type}' 기록: {skipped}건 "
              f"(app.py에서 새로 평가할 때 두 항목을 입력하면 다음부터 집계됩니다)")
    return records


def load_roster(roster_path):
    """전체 명단 파일을 읽어 이름 목록으로 반환. 없으면 None."""
    if not roster_path:
        return None
    path = Path(roster_path)
    if not path.exists():
        print(f"[안내] 명단 파일 '{roster_path}'을 찾을 수 없어 '❌ 미제출' 행은 비워둡니다.")
        return None
    return [
        line.strip()
        for line in path.read_text(encoding="utf-8").splitlines()
        if line.strip() and not line.strip().startswith("#")
    ]


def _week_block_header(week, config):
    cols = [f"{week}주차_{label}" for _, label in config["criteria"]]
    cols.append(f"{week}주차_총점")
    if config["has_team_type"]:
        cols.append(f"{week}주차_팀유형")
    cols.append(f"{week}주차_평가")
    return cols


def _week_block_values(rec, config):
    if not rec:
        width = len(config["criteria"]) + 2 + (1 if config["has_team_type"] else 0)
        return [""] * width
    result = rec["result"]
    vals = [result.get(key, {}).get("score", "") for key, _ in config["criteria"]]
    vals.append(result.get("총점", ""))
    if config["has_team_type"]:
        vals.append(result.get("team_type", ""))
    vals.append(result.get("종합평가", ""))
    return vals


def build_status_sheet(wb, records, report_type):
    """"보고서 제출 현황" 탭: 개인 보고서는 행=학회원(+팀), 팀 Discussion 등은 행=팀."""
    config = REPORT_CONFIGS[report_type]
    entity = config["entity"]
    ws = wb.active
    ws.title = "보고서 제출 현황"

    weeks = sorted({r["week"] for r in records})

    by_key = defaultdict(lambda: {"team": "", "weeks": {}})
    for r in records:
        entry = by_key[r["_entity_key"]]
        if entity == "member":
            team = _team_name_of(r)
            if team:
                entry["team"] = team  # 가장 최근 제출의 팀명을 대표 팀으로 사용
        entry["weeks"][r["week"]] = r

    header = ["학회원", "팀"] if entity == "member" else ["팀"]
    for week in weeks:
        header += _week_block_header(week, config)
    ws.append(header)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for key in sorted(by_key):
        entry = by_key[key]
        row = [key, entry["team"]] if entity == "member" else [key]
        for week in weeks:
            row += _week_block_values(entry["weeks"].get(week), config)
        ws.append(row)

    _autosize(ws)
    ws.freeze_panes = "C2" if entity == "member" else "B2"


def build_weekly_summary_sheet(wb, records, roster, report_type):
    """"주차별 요약" 탭: 순위 Top3 / 제출 / 미제출."""
    config = REPORT_CONFIGS[report_type]
    entity = config["entity"]
    ws = wb.create_sheet("주차별 요약")

    weeks = sorted({r["week"] for r in records})
    by_week = defaultdict(list)
    for r in records:
        by_week[r["week"]].append(r)

    header = ["구분"] + [f"{week}주차" for week in weeks]
    ws.append(header)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL

    rank_labels = ["🏆 1위", "🏆 2위", "🏆 3위"]
    rank_rows = {label: [""] * len(weeks) for label in rank_labels}
    submitted_row = [""] * len(weeks)
    missing_row = [""] * len(weeks)

    for col_idx, week in enumerate(weeks):
        week_records = sorted(by_week[week], key=lambda r: r["result"].get("총점", 0), reverse=True)

        for rank_idx, label in enumerate(rank_labels):
            if rank_idx >= len(week_records):
                continue
            rec = week_records[rank_idx]
            result = rec["result"]
            scores = "/".join(str(result.get(key, {}).get("score", "-")) for key, _ in config["criteria"])
            if entity == "member":
                team = _team_name_of(rec) or "(팀 미입력)"
                rank_rows[label][col_idx] = f"{team}: {rec['_entity_key']} [{scores}]"
            else:
                rank_rows[label][col_idx] = f"{rec['_entity_key']} [{scores}]"

        submitted_row[col_idx] = ", ".join(sorted(r["_entity_key"] for r in week_records))

        if roster:
            submitted_keys = {r["_entity_key"] for r in week_records}
            missing = sorted(set(roster) - submitted_keys)
            missing_row[col_idx] = ", ".join(missing) if missing else "(전원 제출)"
        else:
            missing_row[col_idx] = "(명단 파일 없음)"

    for label in rank_labels:
        ws.append([label] + rank_rows[label])
    ws.append(["✅ 제출"] + submitted_row)
    ws.append(["❌ 미제출"] + missing_row)

    _autosize(ws)


def _autosize(ws, max_width=40):
    for col_cells in ws.columns:
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(length + 2, max_width)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--type", default="개인 보고서", choices=list(REPORT_CONFIGS),
                         help="집계할 보고서 종류 (기본값: 개인 보고서)")
    parser.add_argument("--roster", default=None,
                         help="전체 명단 파일 경로 (미제출자 계산용, 선택. --type에 맞는 이름 단위여야 함)")
    parser.add_argument("--output", default=None, help="출력 엑셀 파일 경로 (기본값: results/제출현황_<종류>.xlsx)")
    args = parser.parse_args()

    report_type = args.type
    records = load_records(report_type)
    if not records:
        print(f"집계할 '{report_type}' 기록이 없습니다 (이름/주차가 채워진 기록이 필요합니다).")
        return

    roster = load_roster(args.roster)

    wb = Workbook()
    build_status_sheet(wb, records, report_type)
    build_weekly_summary_sheet(wb, records, roster, report_type)

    output_path = Path(args.output) if args.output else RESULTS_DIR / f"제출현황_{report_type.replace(' ', '')}.xlsx"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"엑셀 저장 완료: {output_path} ({len(records)}건, {len({r['week'] for r in records})}개 주차)")


if __name__ == "__main__":
    main()
